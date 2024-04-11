
# dependencies
from openpyxl import load_workbook
from openpyxl.utils import get_column_interval
from openpyxl.utils.cell import coordinate_from_string as cfs
import pandas as pd
import os
import warnings


# openpyxl doesn't support the excel data validation fields. ignore the warninig
warnings.filterwarnings('ignore', category = UserWarning, module = 'openpyxl')


# generate data frame
def convert_coordinates_to_dataframe(table_coordinates, sheet):
    col_start = cfs(table_coordinates.split(':')[0])[0]
    col_end = cfs(table_coordinates.split(':')[1])[0]

    # store table contents in array
    data_rows = []
    for row in sheet[table_coordinates]:
        data_rows.append([str(cell.value).replace(",", " ") for cell in row])
    
    # create dataframe from array
    df = pd.DataFrame(data_rows, columns = get_column_interval(col_start, col_end))

    df.columns = df.iloc[0] # Change header to first row
    df = df[1:]  # remove first row from DataFrame to remove the duplicate

    # df.insert(0,'report_exec_date')

    return df


# extract file data
def extract_file(file_path_in, file_path_out):

    # get file name from path
    base = os.path.basename(file_path_in)
    file_name = os.path.splitext(base)[0]
    file_name = file_name.lower()

    # get parent directory
    dirname = os.path.basename(os.path.dirname(file_path_in)) 
    dirname = dirname.lower()

    # variables
    wb              = load_workbook(file_path_in, data_only = True) # workbook (file)
    df_dict         = {} # timesheet hours
    ws_timesheet    = wb['CSM Template'] # worksheet (timesheet)
    ws_user         = wb['Settings'] # worksheet (user details)
    timesheet_data  = pd.DataFrame()
    user_data       = pd.DataFrame()
    selected_user   = pd.DataFrame()
    department      = pd.DataFrame()


    # tables that contain the data that we are interested in
    timesheet_table_list = ['log']
    user_table_list = ['staff', 'member', 'department']


    # get the table coordinates from the worksheet table dictionary
    for tblname, tblcoord in ws_timesheet.tables.items():
        if tblname in timesheet_table_list:
            # print(f"Table Name: {tblname}, Coordinate: {tblcoord}\n")
            df_dict[tblname] = convert_coordinates_to_dataframe(tblcoord, ws_timesheet)  # Convert to dataframe


    for tblname, tblcoord in ws_user.tables.items():
        if tblname in user_table_list:
            # print(f"Table Name: {tblname}, Coordinate: {tblcoord}\n")
            df_dict[tblname] = convert_coordinates_to_dataframe(tblcoord, ws_user)  # Convert to dataframe


    # add spacer in output
    # print("\n")


    # get user name
    selected_user = df_dict['member']
    selected_user_name = (selected_user.to_dict())['Selected Team Member'][1]

    # get department
    department = df_dict['department']
    department['id'] = 1

    # get time data
    timesheet_data = df_dict['log']
    timesheet_data['id'] = 1

    # get user details
    user_data = df_dict['staff']
    user_data.set_index("Name", inplace = True, drop = False) 
    user_data = user_data.loc[[selected_user_name]]
    user_data['id'] = 1

    # prep output
    user_data = user_data.merge(department, on = 'id')
    output = user_data.merge(timesheet_data, on = 'id')
    output = output.drop('id', axis = 1)
    output = output.replace(to_replace = "None", value = "0")

    # Using drop() function to delete last row
    output.drop(index = output.index[-1], axis = 0, inplace = True)

    '''
    # make sure both versions of the template contain the same number of columns
    if 'Vertical' not in output:
        # Insert after VTS
        # print('No vertical')
        output.insert(4, 'Vertical', '')


    if 'Team' not in output:
        # insert after Vertical
        # print('No team')
        output.insert(5, 'Team', '')
    '''

    # print(output)
    new_file_path_out = fr"{file_path_out}/{dirname}/"
    pathExists = os.path.exists(new_file_path_out)
    if not pathExists:
        os.makedirs(new_file_path_out)

    output.to_csv(fr"{new_file_path_out}/{file_name}.csv", header = True, index = False)


